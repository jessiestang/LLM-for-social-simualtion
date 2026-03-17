from openai import OpenAI
from docx import Document
import sys
import json
import os
import re
from pathlib import Path


class ModelConstructor:
    def __init__(self, model_name="gpt-4o-mini", logger = None):
        self.client = OpenAI(
            api_key=os.getenv("OPENAI_API_KEY")
        )  # set your API key in environment variable
        self.model_name = model_name  # default model
        self.logger = logger  # initialize the logger
    
    def _safe_json_load(self, text: str):
        """
        Robustly try to parse JSON from LLM output.
        Strategies:
        1) Direct `json.loads`.
        2) Extract the first {...} or [...] block and parse.
        3) Escape backslashes in the extracted block and parse (helps with LaTeX like `\beta`).
        Raises ValueError with the raw text when parsing ultimately fails.
        """
        try:
            return json.loads(text)
        except json.JSONDecodeError:
            pass

        # Extract first JSON object/array block
        match = re.search(r"(\{.*\}|\[.*\])", text, re.DOTALL)
        if match:
            candidate = match.group(0)
            try:
                return json.loads(candidate)
            except json.JSONDecodeError:
                # Try escaping single backslashes which commonly break JSON when LLM outputs LaTeX
                escaped = candidate.replace("\\", "\\\\")
                try:
                    return json.loads(escaped)
                except json.JSONDecodeError:
                    pass

        raise ValueError(f"LLM output not valid JSON. Raw output:\n{text}")
    
    def _ensure_dict(self, val):
        if isinstance(val, dict):
            return val
        if isinstance(val, list):
            return val
        if isinstance(val, str):
            # Strip markdown fences if present
            cleaned = re.sub(r"```json|```", "", val).strip()
            try:
                return json.loads(cleaned)
            except json.JSONDecodeError:
                # Handle double-escaped strings
                try:
                    return json.loads(cleaned.encode().decode('unicode_escape'))
                except Exception:
                    raise ValueError(f"Cannot parse as JSON: {cleaned[:200]}")
        raise ValueError(f"Expected dict or JSON string, got {type(val).__name__}: {str(val)[:200]}")
    
    def decision_rule_variables(self, problem_context:str):
        """
        This LLM agent takes the problem context as input and think about what variables might by important in driveing the agent's baheviour change."""
        with open(problem_context, "r", encoding="utf-8") as file:
            problem_definition = file.read()
        
        system_prompt = """
        You are a computational social scientist specializing in agent-based modeling (ABM).
        Given the problem context, think about what variables or factors might be important in driving the agent's behavior change, if you were the agent in the model.
        Output a list of potential variables that could influence the agent's decision-making process, along with a brief explanation of why each variable might be relevant.
        Output strictly as *valid JSON* with the following schema:
        Do not give more than 5 variables.
        Check if the variables you proposed are logically consistent with each other and with the problem context.
        {
            "potential_variables": [
                {
                    "name": "variable_name",
                    "explanation": "brief explanation of why this variable is relevant, and how it will affect the decision-making process"
                    "data_type": "data type (e.g., float, integer, boolean)",
                    "extreme_values": {
                        "min": "minimum value",
                        "representation_of_min": "what this minimum value represents in the real world",
                        "max": "maximum value",
                        "representation_of_max": "what this maximum value represents in the real world"
                    }
                }
            ]
        }
        """
        user_prompt = f"""The provided research context is {problem_definition}. Stick strictly to the requirement in the system prompt."""
        response = self.client.chat.completions.create(
            model=self.model_name,
            messages=[
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": user_prompt},
            ],
        )
        pb = response.choices[0].message.content
        variables = self._safe_json_load(pb)
        return variables
    
    def decision_rule_designer(self, problem_context:str, variables:dict):
        """
        This LLM agent will take a list of variables and the problem context as input.
        It will think about how these variables interact together to drive the agent's decision-making process.
        It will express the relationship in if-then rules"""
        variables = self._ensure_dict(variables)
        with open(problem_context, "r", encoding="utf-8") as file:
            problem_definition = file.read()
        
        system_prompt = """
        You are a computational social scientist specializing in agent-based modeling (ABM).
        You will be provided with a problem context and a list of variables relevant to the agent's decision-making process.
        You will also receive user requirements regarding to which variables they wish to include in the decision rule design. Follow strictly the user requirements.

        Now let's do this task step-by-step:
        (1) interaction analysis
            - Reason explicitly about how selected variables interact.
            - Identify if any variables are so closely related that they are better represented as a single composite variable. If so, propose the merge and explain why.
            - Ask yourself:
                * Which variables always move together? → candidate for merging
                * Which variables offset each other? → candidate for a ratio or difference variable
                * Is there a higher-order concept that better captures the agent's mental state?

        (2) rule derivation
        
        Rules must follow this two-part structure:

        Part 1 — Signal composition:
        Compute one or more intermediate signals by combining variables.
        This makes explicit how variables are weighted and aggregated before a decision is made.
        Use named parameters (e.g. beta, weight, threshold) instead of hardcoded numbers.
        No real formulas — express combination qualitatively but in code-like syntax.

        Part 2 - Decision logic:
        Use the computed signal(s) to make the binary decision.
        May have multiple elif branches if there are override conditions

        - Part 1 must reference only variables defined in selected_variables or composite_variables.
        - Part 2 must reference only signals computed in Part 1, or variables with clear ordinal meaning.
        - Named parameters (beta, threshold, etc.) must appear in the "parameters" field.
        - Each behavioral decision still maps to EXACTLY ONE rule block (Part 1 + Part 2 together).

        (3) consistency check
             - Verify that rules do not contradict each other.
             - Verify that every variable selected in step (1) appears in at least one rule.
            - Verify that every behavioral decision in the problem context is covered by exactly one rule.
        
        Here is an example to illustrate the expected output quality:

            Pedestrian Evacuation
            Context: Agents decide whether to evacuate immediately or wait during a building fire.

            Selected variables: threat_proximity, exit_familiarity, crowd_density, panic_level

            Composite variable introduced:
            {
                "name": "evacuation_urgency",
                "composed_from": ["threat_proximity", "crowd_density"],
                "conceptual_meaning": "the combined pressure an agent feels to act immediately,
                                    accounting for both physical danger and social congestion",
                "composition_logic": "high threat_proximity amplified by high crowd_density
                                    produces urgency; if either is low, urgency is dampened",
                "data_type": "float",
                "update_rule": "updated each timestep based on current threat_proximity
                                and observed crowd_density in neighbouring cells"
            }

            Decision rule:
            {
                "behavioral_decision": "evacuate now or wait",
                "outcome_variable": "agent.is_evacuating",
                "signal_computation": [
                    {
                        "signal_name": "evacuation_urgency",
                        "composed_from": ["threat_proximity", "crowd_density"],
                        "expression": "evacuation_urgency = alpha * threat_proximity + (1 - alpha) * crowd_density",
                        "parameter": "alpha — controls relative weight of physical threat vs social congestion"
                    },
                    {
                        "signal_name": "route_confidence",
                        "composed_from": ["exit_familiarity", "panic_level"],
                        "expression": "route_confidence = exit_familiarity * (1 - panic_level)",
                        "parameter": "none — panic directly suppresses familiarity-based confidence"
                    }
                ],
                "rule_pseudocode": 
                    "IF evacuation_urgency > urgency_threshold:
                        agent.is_evacuating = True
                    ELIF route_confidence > confidence_override_threshold:
                        agent.is_evacuating = True   # override: knows the exit well enough to act despite low urgency
                    ELSE:
                        agent.is_evacuating = False",
                "parameters": [
                    {
                        "name": "alpha",
                        "role": "balances physical threat vs crowd pressure in urgency signal",
                        "default": "0.6",
                        "calibratable": true
                    },
                    {
                        "name": "urgency_threshold",
                        "role": "minimum urgency level required to trigger evacuation",
                        "default": "0.5",
                        "calibratable": true
                    },
                    {
                        "name": "confidence_override_threshold",
                        "role": "route confidence level above which agent evacuates regardless of urgency",
                        "default": "0.7",
                        "calibratable": true
                    }
                ],
                "variables_used": ["threat_proximity", "crowd_density", "exit_familiarity", "panic_level"],
                "explanation": "The agent first computes how urgent the situation feels (evacuation_urgency)
                                and how confidently they can navigate to the exit (route_confidence).
                                Evacuation is triggered either when urgency crosses a threshold, OR when
                                the agent knows the route well enough to act despite low perceived urgency.
                                panic_level acts as a suppressor on route confidence — a panicking agent
                                cannot effectively use their spatial knowledge."
            }
 
        Output schema:
        {
            "selected_variables": [
                {
                    "name": "variable_name"
                }
            ],
            "composite_variables": [
                {
                    "name": "composite_variable_name",
                    "composed_from": ["var1", "var2"],
                    "conceptual_meaning": "what this composite represents behaviorally",
                    "composition_logic": "how the components combine — qualitative description, no formula",
                    "data_type": "...",
                    "pseudocode": "...",
                    "update_rule": "..."
                }
            ],
            "interaction_analysis": [
                {
                    "behavioral_decision": "...",
                    "reasoning": "...",
                    "merge_decisions": "explain any merges made and why"
                }
            ],
            "decision_rules": [
                {
                    "behavioral_decision": "...",
                    "outcome_variable": "...",
                    "signal_computation": [
                        {
                            "signal_name": "name of the intermediate signal",
                            "composed_from": ["var1", "var2"],
                            "expression": "how the signal is computed from its components, in code-like pseudocode",
                            "parameter": "param — what it controls (e.g. balances global vs local influence)"
                        }
                    ],
                    "rule_pseudocode": "expression of the decision rule in pseudocode, using variable names, signals, parameters, and thresholds as needed",
                    "parameters": [
                        {
                            "name": "parameter_name",
                            "role": "what it controls in the decision",
                            "default": "theory-grounded default value",
                            "calibratable": true
                        }
                    ],
                    "variables_used": ["list of all variables and signals appearing in the rule"],
                    "explanation": "..."
                }
            ]
        }
        """
        user_prompt = f"""The provided research context is {problem_definition}, and here are the relevant variables: {json.dumps(variables, indent=2)}.
        Stick strictly to the requirement in the system prompt."""
        
        response = self.client.chat.completions.create(
            model=self.model_name,
            messages=[
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": user_prompt},
            ],
        )
        pb = response.choices[0].message.content
        decision_rules = self._safe_json_load(pb)
        return decision_rules
    
    def mechanism_translation(self, problem_context:str, variables:dict, decision_rules:dict, model_save_path:str):
        """
        This LLM agent will regorganize and export a complete and descriptive mechanistic model
        based on the problem context, the preliminary decision-making part, and the secondar decision-making part,
        if available.
        """
        variables = self._ensure_dict(variables)
        decision_rules = self._ensure_dict(decision_rules)
        with open(problem_context, "r", encoding="utf-8") as file:
            problem_definition = file.read()
        
        
        system_prompt="""
        You are a computational social scientist specializing in agent-based modeling (ABM).
        Your task is to transform short, partially specified model notes into a complete,
        descriptive, and formal mechanistic model design suitable for implementation and analysis.

        You will be provided with:
        (1) A problem context file describing the research question, agent attributes, and environmental setup;
        (2) A JSON file describing variables that drive agent behavior;
        (3) A JSON file describing the if-then decision rules that govern agent behavior.


        Your goals are to:
        - Produce a complete and structured model specification that can be directly implemented in code, ensuring that all necessary components for implementation are included;
        - Ensure that no new information or assumptions are introduced beyond what appears in the provided files.

        Follow these steps carefully:

        1. Identify all agent types, their attributes, and actions implied by the rules.
        2. Define the environment (e.g., grid, network) and its parameters.
        3. Describe the behavioural rule and the decision-making process in detail, including the if-then rules, and how variables are updated.
        4. Summarize model-level mechanisms such as feedback loops or external influences.
        5. List all simulation parameters, including agent population size, time steps, constants, and sensitivity parameters.
        6. Make sure that all variables and parameters have a reasonable default value and range (or domain).
        6. Provide a short academic-style description (4-5 sentences) summarizing the full model design.

        Output strictly as *valid JSON* (no Markdown, no commentary).
        Use the following schema exactly:

        {
            "model_title": "short descriptive title",
            "overview": "brief purpose of the model",
            "agents": {
                "types": ["AgentType1", "AgentType2"],
                "attributes": {
                    "AgentType1": [{
                    "attr1": "update rule, datatype and range", 
                    "attr2": "definition and update rule, datatype and range"
                    }],
                    "AgentType2": [{
                    "attr1": "definition and update rule, datatype and range", 
                    "attr2": "definition and update rule, datatype and range"
                    }]
                },
                "actions": {
                    "AgentType1": ["action1", "action2"],
                    "AgentType2": ["action1", "action2"]
                }
            },
            "environment": {
                "structure": "description of environment",
                "interaction_rules": "how agents interact with neighbours or the environment",
                "parameters": {
                    "param1": "description of parameter 1",
                    "default_value_variable_name_1": "a reasonable default value",
                    "range_variable_name_1": "expected numerical range or domain"
                }
            },
            "model_level_mechanisms": [
                {
                    "system_name": "name of the mechanism",
                    "inputs_from_agents": "what data or messages the system receives from agents",
                    "internal_variables": [
                        {
                            "name": "variable_name",
                            "meaning": "what it represents inside the system",
                            "update_rule": "how it updates over time"
                        }
                    ],
                    "processing_logic": [
                        {
                            "rule": "explicit mathematical or algorithmic rule",
                            "description": "what this rule does and why it matters"
                        }
                    ],
                    "outputs_to_agents": [
                        {
                            "output_variable": "name of what is sent back to agents",
                            "generation_rule": "how it is computed from the system's internal state",
                            "timing": "when it becomes available to agents"
                        }
                    ]
                }
            ],
            "decision_rules": [
                {
                    "behavioral_decision": "...",
                    "outcome_variable": "...",
                    "signal_computation": [
                        {
                            "signal_name": "name of the intermediate signal",
                            "composed_from": ["var1", "var2"],
                            "expression": "signal_name = param * var1 + (1 - param) * var2",
                            "parameter": "param — what it controls (e.g. balances global vs local influence)"
                        }
                    ],
                    "rule_pseudocode": "IF signal > threshold_param:\n    outcome = True\nELIF override_var > override_threshold:\n    outcome = True\nELSE:\n    outcome = False",
                    "parameters": [
                        {
                            "name": "parameter_name",
                            "role": "what it controls in the decision",
                            "default": "theory-grounded default value",
                            "calibratable": true
                        }
                    ]
                }
            ],
            "simulation_parameters": {
                "num_agents": "number of agents in the model",
                "time_steps": "number of iterations to simulate",
                "constants": {
                    "constant_name_1": "description of what this constant does",
                    "default_value_constant_name_1": "a reasonable default value",
                    "range_constant_name_1": "expected numerical range or domain"
                }
            },
            "simulation_schedule": [
                "Step 1: description of what happens in this step",
                "Step 2: description of what happens in this step",
                "Step 3: description of what happens in this step",
                "..."
            ]
        }
        """
        user_prompt = f"""Here is the input problem context: {problem_definition},
        here are the variables that drive agent behavior: {json.dumps(variables, indent=2)}, 
        and here are the if-then decision rules that govern agent behavior: {json.dumps(decision_rules, indent=2)}.    
        Stick strictly to the instructions from the system prompt"""
        
        # call the LLM model
        response = self.client.chat.completions.create(
            model=self.model_name,
            messages=[
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": user_prompt},
            ],
        )
        pb = response.choices[0].message.content
        mechanistic_model = self._safe_json_load(pb)
        # save the mchanistic model
        with open(model_save_path, "w", encoding="utf-8") as f:
            json.dump(mechanistic_model, f, indent=2, ensure_ascii=False)

        return mechanistic_model, f"Model saved successfully at {model_save_path}"

    def ODD_formatter(self, problem_context:str, model_save_path:str, file_path:str):
        """
        This function will format the model description into ODD format, which is a standard format for describing agent-based models.
        """
        with open(model_save_path, "r", encoding="utf-8") as file:
            mechanistic_model = json.load(file)
        with open(problem_context, "r", encoding="utf-8") as file:
            problem_definition = file.read()
        
        system_prompt = """
        You are a computational social scientist specializing in agent-based modeling (ABM).
        Your task is to convert a detailed mechanistic model description into the ODD (Overview, Design concepts, Details) format, which is a standard for describing agent-based models.
        
        You will be provided with:
        (1) A problem context file describing the research question, agent attributes, and environmental setup;
        (2) A JSON file describing the agent decision context, behavioral rules, variables, thresholds, and model parameters.
        
        Your goals are to:
        - Produce a complete and structured model specification in ODD format, ensuring that all necessary components for implementation are included;
        - Ensure that no new information or assumptions are introduced beyond what appears in the provided files.

        Your output should strictly follow the ODD format, which includes the following sections:
        1. Purpose: a concise statement of the model's overall objective and research question.
        2. Entities, state variables, and scales: a detailed description of the agents, their attributes, the environment, and the temporal and spatial scales of the model.
        3. Process overview and scheduling: a step-by-step outline of the processes that occur in each time step, including the order of agent actions and interactions. When and how are state variables updated?
        4. Design concepts: an explanation of the key design principles and mechanisms that drive the model, such as emergence, adaptation, learning, and interaction patterns.
            a.Emergence: Which key model results or outputs are modeled as emerging from the adaptive decisions and behaviors of agents.
            b.Adaptation: How do agents adapt their behavior based on their experiences or changes in the environment?
            c.Learning: Do agents learn from their interactions or outcomes? If so, how is this learning process modeled?
            d.Objectives: What are the goals or objectives that guide agent behavior? Are they maximizing utility, following heuristics, or something else?
            e.Prediction: if an agent's adaptive traits or learning procedures are based on estimating future consequences of decisions, how do agents predict the future conditions (either environmental or internal) they will experience?
            f.Sensing: What information do agents have access to when making decisions? Do they have perfect information about their environment, or do they rely on local perceptions or heuristics?
            g.Interaction: How do agents interact with each other and with the environment? Are interactions local or global, and what is the nature of these interactions (e.g., competition, cooperation, communication)?
            h.Stochasticity: What role does randomness play in the model? Are there stochastic elements in agent decision-making, interactions, or environmental changes?
            i.Collectives: Do the individuals form or belong to aggregations that affect, and are affected by, the individuals?
            j.Observation: What data or outputs are collected from the model, and how do they relate to the research question?
        5. Initialization: a description of how the model is initialized (t = 0), including the initial conditions of agents and the environment.
        6. Input data: a description of any external data that is used to drive the model (if there is), including how it is incorporated and its role in the model dynamics.
        7. Submodels: a detailed description of the submodels that govern specific processes or behaviors in the model, including any mathematical equations or decision rules (if there is any).

        An example of the expected output format can be:
        "Specifically, we are addressing the following questions: [purpose]. The model includes the following entities [entities]. 
        They are characterized by the following state variables [state variables]. 
        The spatial and temporal resolution and extent are [temporal and spatial scales]. 
        The most important design concepts of the model are [all relevant items inthe design concepts].
        The model is initialized with [initialization]. 
        Model dynamics are driven by input data representing [input data description]. (only if there is input data)
        We also include the following submodels to capture key processes in the system: [submodel overview]. (only if there is submodel)
        "
        Export the model description in academic text format that follows the ODD template, ensuring that each section is clearly labeled and contains the relevant information extracted from the provided files.
        """

        
        user_prompt = f"""Here is the input problem context: {problem_definition},
        and here is the JSON file describing the agents'behavior and decision-making {mechanistic_model}.
        Stick strictly to the instructions from the system prompt"""
        
        # call the LLM model
        response = self.client.chat.completions.create(
            model=self.model_name,
            messages=[
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": user_prompt},
            ],
        )
        text = response.choices[0].message.content

        # save the ODD text to a Word file
        doc = Document()
        doc.add_heading("ODD Model Description", level=1)

        # Split into lines
        lines = text.split("\n")

        for line in lines:
            line = line.strip()

            if not line:
                doc.add_paragraph("")  # preserve empty line
                continue
            doc.add_paragraph(line)
        doc.save(file_path)

        return text,f"ODD files save successfully to {file_path}"
    

    """
    You are a computational social scientist specializing in agent-based modeling (ABM).
        You will be provided with a problem context and a list of variables relevant to the agent's decision-making process.
        You will also receive user requirements regarding to which variables they wish to include in the decision rule design. Follow strictly the user requirements.

        Let's do this step-by-step:
        Firstly, for each variable in the list:
        (a) Describe in plain language how an INCREASE in this variable affects the agent's behavior.
        (b) Describe how a DECREASE does the same

        Secondly, based on the first step, try to describe the relationship between variables:
        (a) Which variables always move together? → candidate for merging into a composite
        (b) Which variables oppose each other? → candidate for a ratio or difference 
        (c) Which variable has a stonger influence on the agent's behavior? → candidate for a weighted variable
        For each composite variable you propose, write pseudocode showing how it is constructed and explain WHY this composite is more meaningful than its components separately.
        If you include a parameter, also explain why it is needed, and the range of this parameter.

        Thirdly, based on the previous two steps, write up the decision-logic rule for each behaviour in the problem context. DO NOT write separate rules for binary outcomes of one behaviour.
        Do this in the form of pseudocode, in if-else statement.
        You need to use all the variables (including the composite one in the final rule)
        If you include a parameter, also explain why it is needed, and the range of this parameter.
        Describe in word how this rule corresponds to a real-life behaviour.

        Constrain:
        Do not include any hard-coded values in any rules; replace them with proper variable names

        Output strictly as *valid JSON* with the following schema:
        {
            "variables": [
                {
                    "name": "change_variable_name",
                    "change_direction": "increase or decrease",
                    "affected_variable": "description of which variables are affected",
                    "affected_behavior": "description of how agents'behavoir is affected"
                }
            ],
            "composite_variables": [
                {
                    "name": "composite_variable_name",
                    "composed_from": ["var1", "var2"],
                    "conceptual_meaning": "what this composite represents behaviorally",
                    "composition_logic": "how the components combine — qualitative description, no formula",
                    "data_type": "...",
                    "pseudocode": "...",
                    "update_rule": "..."
                }
            ],
            "interaction_analysis": [
                {
                    "variables_move_together": "Description of variables that move in the same direction",
                    "variables_oppose_each_other": "Description of variables that oppose each other",
                    "variables_independent": "Description of variables that are independent"
                }
            ],
            "decision_rules": [
                {
                    "behavioral_decision": "...",
                    "outcome_variable": "...",
                    "rule_pseudocode": "...",
                    "parameters": [
                        {
                            "name": "parameter_name",
                            "role": "what it controls in the decision",
                            "default": "theory-grounded default value",
                            "range": "expected range of the value"
                        }
                    ],
                    "explanation": "How the rule captures real-life behavior"
                }
            ]
        }
    """

    """
    You are a computational social scientist specializing in agent-based modeling (ABM).
        You will be provided with a problem context and a list of variables relevant to the agent's decision-making process.
        You will also receive user requirements regarding to which variables users wish to include in the decision rule design. Follow strictly the user requirements.

        Let's do this step-by-step:
        Imagine you are the agent in the model, and you are making a behavourial decision based on the provided variables.
        Firstly, for each variable in the list:
        - Describe in plain language how change in values (in both directions) of this variables may have an influence on your behavouir.
        E.g., "An increase in variable X makes me more likely to do Y because... A decrease in variable X makes me less likely to do Y because..."'
        - Describe which variables have the most influence on your behavior, and which variables have a weaker influence. In other words, what weight you would give them.
        E.g., "Variable X has the strongest influence on my behavior, because..., so I would give it a weight of [number between 0 and 1]. Variable Y has a weaker influence on my behavior, because..., so I would give it a weight of [number between 0 and 1]..." 
        - Now try to summarize your answers to the above questions into a if-else style decision rule. Make sure you use all the variables, including the weight and threhold variables.
        - Output your decision rule in the form of pseudocode, for example:
        IF variable1 * weight1 + variable2 * weight2 > threshold:
        THEN do Action A
        ELSE do Action B
        - Describe in words how this decision rule captures the real-life behavior.

        Constrain:
        A decision signal must use AT MOST 2 variables or composites.
        - If you have more than 2 variables, you MUST first merge related variables 
        into a composite in Step 2 before writing the rule.
        - A rule that sums all variables with weights is NOT acceptable — 
        it means you skipped the composite variable step.

        BAD (not acceptable):
        signal = var1 * w1 + var2 * w2 + var3 * w3 + var4 * w4
        IF signal > threshold: speak
        GOOD (acceptable):
        social_pressure = var1 * w1 + (1 - w1) * var2   # composite of related variables
        IF social_pressure > threshold AND var3 > var3_threshold:
            speak
        ELIF var4 > var4_threshold:
            speak                                      
        ELSE:
            silent

        Output strictly as valid JSON with the following schema:
        {
            "variable_analysis": [
                {
                    "variable_name": "name of the variable",
                    "influence_analysis": "your plain language analysis of how value changes in both directions affect behavior and other variables",
                    "weight_reasoning": "your reasoning about the weight of this variable relative to others, including the assigned weight",
                    "threshold_reasoning": "your reasoning about whether a threshold exists, and if so, what it is and why"
                }
            ],
            "decision_rules": [
                {
                    "behavioral_decision": "name of the behavior being decided",
                    "reasoning": "your plain language summary of how the variables combine to drive this decision",
                    "rule_pseudocode": "express the decision rule in pseudocode, using variable names, weights, and thresholds as needed",
                    "real_life_interpretation": "one sentence describing how this rule captures real-life behavior"
                }
            ]
        }"""
    

    
    """Imagine you are an agent in the model, and you are making a behavourial decision based on the provided variables.
        Firstly, think about 3-4 scenarios where the variables have different values.
        The variables need to cover a wide range of possible combinations of variable values, including edge cases (e.g., all variables are high, all variables are low, some variables are high while others are low, etc.).
        Describe these scenarios in plain language, and what actions you would take in each scenario.
        Describe some general rules or patterns you see in how the variables influence your behavior across these scenarios.
        Summarize these rules in the form of if-else style decision rules. Make sure you use all the variables, and do not include any hard-coded values in any rules; replace them with proper variable names.
        Output strictly as *valid JSON* with the following schema:
        {
            "scenarios": [
                {
                    "scenario_id": "1",
                    "description": "plain language description of the situation",
                    "variable_values": {
                        "variable_name_1": "high / low / moderate — with brief justification",
                        "variable_name_2": "high / low / moderate — with brief justification"
                    },
                    "action_taken": "the action the agent takes in this scenario",
                    "reasoning": "one sentence explaining why this action maximizes payoff"
                }
            ],
            "patterns": [
                {
                    "observation": "a general pattern noticed across scenarios",
                    "supporting_scenarios": ["list of scenario_ids that support this observation"],
                }
            ],
            "decision_rules": [
                {
                    "behavioral_decision": "name of the behavior being decided",
                    "rule_pseudocode": "express the decision rule in pseudocode, using variable names and thresholds as needed",
                    "parameters": [
                        {
                            "name": "parameter_name",
                            "role": "what it controls in the decision",
                            "range": "expected numerical range e.g. 0.0 to 1.0",
                            "default": "theory-grounded default value"
                        }
                    ],
                    "real_life_interpretation": "one sentence describing what this rule captures in real life"
                }
            ]
        }
        This LLM agent will take a list of variables and the problem context as input.
        It will think about how these variables interact together to drive the agent's decision-making process.
        It will express the relationship in if-then rules"""
    


        